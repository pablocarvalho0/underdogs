"""Parser do relatório da Ponte Câmbio a partir do e-mail .eml recebido."""

import email
import io
from email import policy
from pathlib import Path
from typing import Any

import openpyxl
import pandas as pd

from src.config import PONTE_EML
from src.loaders.taxas import clean_cnpj


def load_ponte_cambio_report(file_path: Path = PONTE_EML) -> pd.DataFrame:
    """
    Lê o arquivo .eml da Ponte Câmbio, extrai programaticamente
    o anexo .xlsx e carrega a aba 'Operações'.
    """
    with open(file_path, "rb") as f:
        msg = email.message_from_binary_file(f, policy=policy.default)

    xlsx_bytes: bytes | None = None
    for part in msg.walk():
        filename = part.get_filename()
        if filename and filename.lower().endswith(".xlsx"):
            payload = part.get_payload(decode=True)
            if isinstance(payload, bytes):
                xlsx_bytes = payload
            break

    if xlsx_bytes is None:
        raise FileNotFoundError("Anexo Excel (.xlsx) não encontrado no e-mail da Ponte Câmbio.")

    wb = openpyxl.load_workbook(io.BytesIO(xlsx_bytes), data_only=True)
    sheet = wb["Operações"] if "Operações" in wb.sheetnames else wb.worksheets[0]

    # Localizar linha do cabeçalho ('Nº Operação')
    header_row_idx = None
    headers = []
    for idx, row in enumerate(sheet.iter_rows(values_only=True), start=1):
        row_str = [str(v).strip().lower() if v is not None else "" for v in row]
        if any("operação" in s or "operacao" in s for s in row_str):
            header_row_idx = idx
            headers = [str(v).strip() if v is not None else f"col_{i}" for i, v in enumerate(row)]
            break

    if header_row_idx is None:
        raise ValueError("Cabeçalho da planilha de operações da Ponte Câmbio não encontrado.")

    # Abaixo do cabeçalho vêm as operações e, no fim, linhas de rodapé do tipo
    # ("Total de operações:", 233). O rodapé fica em df.attrs para os invariantes.
    data = []
    rodape: dict[str, Any] = {}
    for linha_planilha, row in enumerate(
        sheet.iter_rows(min_row=header_row_idx + 1, values_only=True), start=header_row_idx + 1
    ):
        if not any(row):
            continue
        primeira = str(row[0]).strip() if row[0] is not None else ""
        if primeira.endswith(":"):
            rodape[primeira.rstrip(":").strip()] = row[1]
            continue
        row_dict = dict(zip(headers, row))
        row_dict["linha_planilha"] = linha_planilha  # distingue cópias com o mesmo Nº Operação
        data.append(row_dict)

    df = pd.DataFrame(data)
    if "CNPJ Cliente" in df.columns:
        df["cnpj_limpo"] = df["CNPJ Cliente"].apply(clean_cnpj)
    df.attrs["rodape"] = rodape

    return df
