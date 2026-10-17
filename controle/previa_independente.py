"""Prévia independente de setembro: recalcula os números de controle direto dos arquivos brutos.

Não importa nada de `src/`: é a régua contra a qual o pipeline é conferido (HIPOTESES.md §5).
Lê o data pack, aplica as regras do HIPOTESES.md e compara com `controle/esperado_setembro.json`.

Uso: python controle/previa_independente.py [pasta_do_data_pack]
Sai com código 1 se algum número não bater.
"""

import collections
import csv
import datetime as dt
import email
import io
import json
import re
import sys
from decimal import ROUND_HALF_UP, Decimal
from email import policy
from pathlib import Path

import openpyxl

RAIZ = Path(__file__).resolve().parent.parent
PASTA = Path(sys.argv[1]) if len(sys.argv) > 1 else RAIZ / "docs" / "Tarvo_Data_Pack_set2026"
ESPERADO = json.loads((RAIZ / "controle" / "esperado_setembro.json").read_text(encoding="utf-8"))

CASA = "00520815000100"  # R-08: tarifa zero no checkout de 01/09 a 31/12/2026
ISENCAO = (dt.date(2026, 9, 1), dt.date(2026, 12, 31))
ALIAS = {"1043300981": "66564512000199", "1043396662": "78576234000179", "1038867654": "07752116000126"}  # R-05
PONTE = "75374179000146"  # R-27
SPREAD_PADRAO = Decimal("0.015")  # R-12
TOLERANCIA = Decimal("0.0001")  # P-01: 0,01 p.p.


def brl(x: Decimal) -> Decimal:  # R-02: meio centavo para cima, por linha
    return x.quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)


def dec_br(s: str) -> Decimal:
    return Decimal(s.replace(".", "").replace(",", ".")) if s.strip() else Decimal(0)


def data_br(s: str) -> dt.date:
    return dt.datetime.strptime(s.strip(), "%d/%m/%Y").date()


def so_digitos(s: object) -> str:
    return re.sub(r"\D", "", str(s)).zfill(14)


def pct(v: object) -> Decimal | None:
    if v is None or v in ("", "padrão"):
        return None
    if isinstance(v, str):
        return Decimal(v.replace("%", "").replace(",", ".")) / 100
    return Decimal(str(v))


def ler_texto(p: Path) -> str:
    try:
        return p.read_text(encoding="utf-8-sig")
    except UnicodeDecodeError:
        return p.read_text(encoding="cp1252")


# --- Planilha de taxas (aba vigente; INATIVO fora) ---
xlsx = next(PASTA.glob("*.xlsx"))
aba = openpyxl.load_workbook(xlsx).worksheets[0]
tabela: dict[str, list] = collections.defaultdict(list)
spread_contrato: dict[str, Decimal] = {}
for r in aba.iter_rows(min_row=6, max_row=40, values_only=True):
    if not r[1]:
        continue
    nome, cnpj_txt, ec, deb, av, p26, p712, pix, ant, spr, vig, obs = r[:12]
    cnpj = so_digitos(cnpj_txt)
    spread_contrato[cnpj] = pct(spr) if pct(spr) is not None else SPREAD_PADRAO
    if obs and "INATIVO" in str(obs):
        continue
    ini, fim = dt.date(2000, 1, 1), dt.date(2100, 1, 1)
    if isinstance(vig, str) and vig.startswith("até"):
        fim = data_br(vig.split()[-1])
    if isinstance(vig, str) and vig.startswith("a partir"):
        ini = data_br(vig.split()[-1])
    tabela[cnpj].append((ini, fim, {"DEBITO": pct(deb), "AV": pct(av), "P26": pct(p26), "P712": pct(p712), "PIX": pct(pix), "ANT": pct(ant)}))
    if ec:
        ALIAS.setdefault(str(ec), cnpj)

# --- Orbe: todo CSV com o cabeçalho da Orbe (01 + 06, R-43); linha idêntica conta uma vez ---
linhas_orbe: list[dict] = []
vistas: set[tuple] = set()
for p in sorted(PASTA.iterdir()):
    if p.suffix.lower() != ".csv":
        continue
    texto = ler_texto(p)
    if not texto.startswith("id_transacao"):
        continue
    for r in csv.DictReader(io.StringIO(texto), delimiter=";"):
        chave = tuple(r.values())
        if chave not in vistas:
            vistas.add(chave)
            linhas_orbe.append(r)
por_id = {r["id_transacao"]: r for r in linhas_orbe}

comissao: dict[str, Decimal] = collections.defaultdict(Decimal)
divergencias: list[dict] = []
casa_devolucao, casa_linhas = Decimal(0), 0
sem_contrato: dict[str, Decimal] = collections.defaultdict(Decimal)
for r in linhas_orbe:
    ec, tipo, mod = r["cod_estabelecimento"], r["tipo_lancamento"], r["modalidade"]
    cnpj = ALIAS.get(ec)
    if cnpj is None or not tabela.get(cnpj):
        sem_contrato[ec] += dec_br(r["valor_taxa"])  # R-06: não fatura
        continue
    if tipo == "ESTORNO":
        base = data_br(r["data_venda_original"])  # R-03
    elif tipo == "REVERSAO ANTECIPACAO":
        base = data_br(por_id[r["id_referencia"]]["data_movimento"])  # R-04
    else:
        base = data_br(r["data_movimento"])
    if tipo in ("ANTECIPACAO", "REVERSAO ANTECIPACAO"):
        faixa = "ANT"
    else:
        faixa = {"DEBITO": "DEBITO", "PIX": "PIX", "CREDITO A VISTA": "AV"}.get(mod) or (
            "P26" if int(r["qtd_parcelas"]) <= 6 else "P712"
        )
    taxa = next(t for ini, fim, t in tabela[cnpj] if ini <= base <= fim)[faixa]
    bruto, aplicada = dec_br(r["valor_bruto"]), dec_br(r["taxa_aplicada_pct"]) / 100
    isento = cnpj == CASA and ISENCAO[0] <= base <= ISENCAO[1]
    comissao[cnpj] += Decimal(0) if isento else brl(bruto * taxa)  # R-01, R-02
    if isento:  # R-09: devolução, não divergência
        casa_devolucao += dec_br(r["valor_taxa"])
        casa_linhas += 1
        continue
    if aplicada != taxa:  # R-20, R-21
        divergencias.append({"linha_negocio": "CHECKOUT", "cnpj": cnpj, "taxa_contratada_pct": f"{(taxa * 100).normalize():f}", "diferenca_brl": brl(bruto * (aplicada - taxa))})

# P-08: lojista sem contrato (novo) cuja antecipação acumulada supera as vendas acumuladas
alerta: dict[str, Decimal] = {}
for ec in sem_contrato:
    dia_a_dia: dict[dt.date, list[Decimal]] = collections.defaultdict(lambda: [Decimal(0), Decimal(0)])
    for r in linhas_orbe:
        if r["cod_estabelecimento"] != ec:
            continue
        d = data_br(r["data_movimento"])
        if r["tipo_lancamento"] == "VENDA":
            dia_a_dia[d][1] += dec_br(r["valor_bruto"])
        elif r["tipo_lancamento"] in ("ANTECIPACAO", "REVERSAO ANTECIPACAO"):
            dia_a_dia[d][0] += dec_br(r["valor_bruto"])
    ant = ven = pico = Decimal(0)
    for d in sorted(dia_a_dia):
        ant += dia_a_dia[d][0]
        ven += dia_a_dia[d][1]
        pico = max(pico, ant - ven)
    if pico > 0:
        alerta[ec] = pico

# --- Ponte: anexo do e-mail ---
eml = next(PASTA.glob("*.eml"))
msg = email.message_from_bytes(eml.read_bytes(), policy=policy.default)
anexo = next(a for a in msg.iter_attachments() if str(a.get_filename() or "").lower().endswith(".xlsx"))
wb = openpyxl.load_workbook(io.BytesIO(anexo.get_payload(decode=True)), data_only=True)
ops_ws = wb["Operações"] if "Operações" in wb.sheetnames else wb.worksheets[0]
cab_linha = next(i for i, row in enumerate(ops_ws.iter_rows(values_only=True), 1) if row and "Nº Operação" in row)
cab = [c.value for c in ops_ws[cab_linha]]
operacoes = [dict(zip(cab, [c.value for c in row])) for row in ops_ws.iter_rows(min_row=cab_linha + 1) if str(row[0].value or "").startswith("PC-")]

remessa: dict[str, Decimal] = collections.defaultdict(Decimal)
negativos: list[Decimal] = []
duplicata = Decimal(0)
vistos: set[str] = set()
for o in operacoes:
    n = o["Nº Operação"]
    spread = Decimal(str(o["Spread R$"]))
    if n in vistos:  # R-13: conta uma vez
        duplicata += spread
        continue
    vistos.add(n)
    cnpj = so_digitos(o["CNPJ Cliente"])
    if spread < 0:  # R-14: fora da comissão e da nota
        negativos.append(spread)
        continue
    remessa[cnpj] += brl(spread / 2)  # R-11
    contratado, aplicado = spread_contrato.get(cnpj, SPREAD_PADRAO), Decimal(str(o["Spread %"]))
    if abs(aplicado - contratado) >= TOLERANCIA:  # R-23, P-01, P-05
        divergencias.append({"linha_negocio": "REMESSA", "cnpj": cnpj, "taxa_contratada_pct": f"{(contratado * 100).normalize():f}", "diferenca_brl": brl((aplicado - contratado) * Decimal(str(o["Valor ME"])) * Decimal(str(o["PTAX Ref."])))})

# --- Extratos ---
aurora_csv = next(p for p in PASTA.iterdir() if p.suffix.lower() == ".csv" and "extrato de conta corrente" in (t := ler_texto(p)[:300].lower()) and "aurora" in t)
aurora = {i: l.split(";") for i, l in enumerate(ler_texto(aurora_csv).splitlines(), 1)}
movimentos_aurora = [c for c in aurora.values() if re.match(r"\d{2}/\d{2}/\d{4}$", c[0]) and len(c) > 3 and c[3].strip()]
ofx = next(PASTA.glob("*.ofx"))
movimentos_meridiano = ler_texto(ofx).upper().count("<STMTTRN>")

# --- Saídas calculadas, no formato que o JSON filtra ---
saida = {
    "comissoes_por_lojista.csv": [{"linha_negocio": "CHECKOUT", "cnpj": c, "comissao_brl": v} for c, v in comissao.items()]
    + [{"linha_negocio": "REMESSA", "cnpj": c, "comissao_brl": v} for c, v in remessa.items()],
    "notas_a_emitir.csv": [{"linha_negocio": "CHECKOUT", "tomador_cnpj": c, "valor_brl": v} for c, v in comissao.items() if v > 0]  # R-26
    + [{"linha_negocio": "REMESSA", "tomador_cnpj": PONTE, "valor_brl": sum(remessa.values(), Decimal(0))}],
    "divergencias_taxa.csv": divergencias,
    "lancamentos_classificados.csv": [{"banco": "AURORA"}] * len(movimentos_aurora) + [{"banco": "MERIDIANO"}] * movimentos_meridiano,
    "pendencias.csv": [{"tipo": "SPREAD_NEGATIVO", "valor_brl": v} for v in negativos],
}


def div_por(cnpj: str) -> Decimal:
    return sum((d["diferenca_brl"] for d in divergencias if d["cnpj"] == cnpj and d["linha_negocio"] == "CHECKOUT"), Decimal(0))


def valor_aurora(n: int) -> Decimal:
    return dec_br(aurora[n][3])


pendencias = {
    "PEN-CHK-SEMCONTRATO-1076801732": sem_contrato.get("1076801732"),
    "PEN-CHK-DEVOLUCAO-00520815000100": casa_devolucao,
    "PEN-CHK-SEMNOTA-00520815000100": comissao[CASA],
    "PEN-CHK-CONTESTACAO-95927868000148": div_por("95927868000148"),
    "PEN-CHK-CONTESTACAO-22224579000109": div_por("22224579000109"),
    "PEN-CHK-ALERTA-ANTECIPACAO-1076801732": alerta.get("1076801732"),
    "PEN-REM-CONTESTACAO-17217888000111": sum((d["diferenca_brl"] for d in divergencias if d["linha_negocio"] == "REMESSA"), Decimal(0)),
    "PEN-REM-DUPLICATA": duplicata,
    "PEN-REM-REPASSE-75374179000146": sum(remessa.values(), Decimal(0)),
    "PEN-EXT-NAOIDENT-L50": valor_aurora(50),
    "PEN-EXT-DEVOLUCAO-L65": valor_aurora(65),
    "PEN-EXT-TRANSF-L92": valor_aurora(92),
}
contagens = {"PEN-CHK-DEVOLUCAO-00520815000100": casa_linhas}

# --- Comparação com o JSON ---
falhas = 0
print(f"Prévia independente de {PASTA.name}\n")
for t in ESPERADO["totais"]:
    linhas = [r for r in saida[t["arquivo"]] if all(str(r.get(k)) == v for k, v in t["filtro"].items())]
    obtido = Decimal(len(linhas)) if t["metrica"] == "linhas" else sum((Decimal(r[t["coluna"]]) for r in linhas), Decimal(0))
    ok = obtido == Decimal(t["esperado"])
    falhas += not ok
    print(f"{'ok ' if ok else 'ERRO'} {t['item']:<55} esperado {t['esperado']:>12}  prévia {obtido:>12}")
print()
for p in ESPERADO["pendencias"]:
    obtido = pendencias.get(p["id_saida"])
    ok = obtido is not None and (p["esperado"] is None or obtido == Decimal(p["esperado"]))
    extra = ""
    if "linhas_rastro" in p:
        n = contagens.get(p["id_saida"])
        ok = ok and n == p["linhas_rastro"]
        extra = f"  linhas {n}"
    falhas += not ok
    print(f"{'ok ' if ok else 'ERRO'} {p['id_saida']:<42} esperado {p['esperado']!s:>12}  prévia {obtido}{extra}")
print(f"\n{'Tudo bate.' if not falhas else f'{falhas} item(ns) não batem.'}")
sys.exit(1 if falhas else 0)
